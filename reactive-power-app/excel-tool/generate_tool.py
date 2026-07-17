#!/usr/bin/env python3
"""Generates Reactive_Power_Compensation_Tool.xlsx — a standalone Excel
version of the reactive power sizing app: full calculation chain with live
formulas, a grid-code requirement database with dropdown selection, an
editable inverter P-Q capability table with bilinear interpolation, and a
compliance sheet with interpolated V-Q / P-Q curves and styled charts.
No macros; plain formulas only."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.utils import get_column_letter

# ---------- palette ----------
HEAD = PatternFill("solid", fgColor="1C5CAB")
SUBHEAD = PatternFill("solid", fgColor="D9E4F2")
INPUT = PatternFill("solid", fgColor="FFF3CC")
CALC = PatternFill("solid", fgColor="EEF1F4")
RESULT = PatternFill("solid", fgColor="E6F4E6")
PASSF = PatternFill("solid", fgColor="C9EFC9")
FAILF = PatternFill("solid", fgColor="F6C9C9")
HEADF = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1C5CAB")
THIN = Border(*[Side(style="thin", color="C9CFD6")] * 4)
EPS = "0.001"


def style(ws, cell, fill=None, bold=False, num=None, center=False):
    c = ws[cell]
    if fill is not None:
        c.fill = fill
    if bold:
        c.font = BOLD
    if num:
        c.number_format = num
    if center:
        c.alignment = Alignment(horizontal="center")
    c.border = THIN
    return c


def section(ws, row, text, span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=text)
    c.font = HEADF
    c.fill = HEAD
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18


def title(ws, text, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.cell(row=1, column=1, value=text).font = TITLE
    ws.row_dimensions[1].height = 24


def table_header(ws, row, headers, start_col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + j, value=h)
        c.font = BOLD
        c.fill = SUBHEAD
        c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def passfail(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=PASSF))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAILF))


wb = openpyxl.Workbook()

# =============================== README ===============================
rd = wb.active
rd.title = "README"
rd.sheet_view.showGridLines = False
rd.column_dimensions["A"].width = 112
title(rd, "REACTIVE POWER COMPENSATION SIZING TOOL", 1)
rows = [
    "",
    "A standalone Excel version of the reactive power sizing app. Sheets:",
    "   1. Sizing — the calculation chain: grid demand, transformer consumption, losses, inverter sizing, checks.",
    "   2. GridCodes — database of published grid-code reactive requirements; pick one from the dropdown (B1).",
    "   3. InverterPQ — the manufacturer's Q(t)-mode capability table (% of rated kVA per voltage and loading).",
    "   4. Compliance — V-Q at Registered Capacity and P-Q at nominal voltage vs the selected grid code, with charts.",
    "",
    "Colour key:   yellow = input cells,   grey = calculated,   green = key results.",
    "",
    "Notes:",
    "   - Inverter count: leave the manual-count cell blank to auto-size (meets both reactive and delivery needs).",
    "   - Reactive capability precedence: capability table (if 'Use table' = Yes) > manual kVAr override > sqrt(S²-P²).",
    "   - Compliance dispatches inverters at (grid limit + losses) / count, as in a G99 Type C study.",
    "   - Export capability is net of transformer consumption; absorption is aided by it; capacitor bank exports only.",
    "   - Voltage capability is interpolated linearly between the capability-table rows (conservative across N/A gaps).",
    "   - Grid-code presets are typical published values — always confirm against your interconnection agreement.",
    "   - This tool is an analytic screening aid, not a substitute for a network load-flow study.",
]
for i, t in enumerate(rows, start=2):
    rd[f"A{i}"] = t
    if t in ("Notes:",):
        rd[f"A{i}"].font = BOLD

# =============================== SIZING ===============================
sz = wb.create_sheet("Sizing")
sz.sheet_view.showGridLines = False
for col, w in zip("ABCDE", (52, 14, 10, 2, 60)):
    sz.column_dimensions[col].width = w
title(sz, "SIZING", 3)


def srow(r, label, value, unit, fill, num=None, bold=False):
    sz[f"A{r}"] = label
    sz[f"A{r}"].border = THIN
    if bold:
        sz[f"A{r}"].font = BOLD
    if value is not None:
        sz[f"B{r}"] = value
    sz[f"C{r}"] = unit
    sz[f"C{r}"].border = THIN
    style(sz, f"B{r}", fill, bold=bold, num=num)


section(sz, 3, "PROJECT", 3)
srow(4, "Project name", "Belvior", "", INPUT)

section(sz, 6, "STEP 1 — INITIAL SIZING", 3)
srow(7, "Grid injection limit / Registered Capacity", 52.8, "MW", INPUT, "0.0")
srow(8, "Power factor at grid", 0.95, "cos φ", INPUT, "0.000")
srow(9, "MVA rating of plant", "=B7/B8", "MVA", CALC, "0.00")
srow(10, "MVAr demand (grid PF)", "=SQRT(B9^2-B7^2)", "MVAr", CALC, "0.00")

section(sz, 12, "STEP 2 — TRANSFORMER REACTIVE CONSUMPTION", 3)
table_header(sz, 13, ["Group", "Qty", "Rating (kVA)", "", "Z (%) | Q (kVAr)"])
sz["B13"] = "Qty"
# custom 5-col table across A..E
table_header(sz, 13, ["Group", "Qty", "kVA", "Z %", "Q (kVAr)"])
xmers = [("TS Block 1", 3, 10164, 9.5), ("TS Block 2", 4, 7260, 9.5),
         ("Auxiliary transformer (MCR)", 1, 500, 4.5), ("", None, None, None),
         ("", None, None, None), ("", None, None, None)]
for i, (n, q, kva, z) in enumerate(xmers):
    r = 14 + i
    sz[f"A{r}"] = n
    if q is not None:
        sz[f"B{r}"], sz[f"C{r}"], sz[f"D{r}"] = q, kva, z
    for col in "ABCD":
        style(sz, f"{col}{r}", INPUT)
    sz[f"E{r}"] = f"=IF(COUNT(B{r}:D{r})=3,B{r}*C{r}*D{r}/100,0)"
    style(sz, f"E{r}", CALC, num="0.0")
srow(20, "Design margin on transformer demand", 10, "%", INPUT, "0")
srow(21, "Transformer demand incl. margin", "=SUM(E14:E19)/1000*(1+B20/100)", "MVAr", CALC, "0.00")
srow(22, "TOTAL REACTIVE DEMAND", "=B10+B21", "MVAr", RESULT, "0.00", bold=True)

section(sz, 24, "STEP 3 — INTERNAL SYSTEM LOSSES", 3)
table_header(sz, 25, ["Loss source", "Fraction %", "Loss (MW)"])
losses = [("AC cable loss", 1.0), ("Auxiliary system loss (string inverters)", 1.0),
          ("Transmission line loss", 0.3), ("Miscellaneous losses", 1.0)]
for i, (n, pct) in enumerate(losses):
    r = 26 + i
    sz[f"A{r}"] = n
    sz[f"A{r}"].border = THIN
    sz[f"B{r}"] = pct
    style(sz, f"B{r}", INPUT, num="0.0")
    sz[f"C{r}"] = f"=B{r}*$B$7/100"
    style(sz, f"C{r}", CALC, num="0.000")
srow(30, "Total system losses", "=SUM(C26:C29)", "MW", CALC, "0.000")
srow(31, "Capacity to be considered", "=B7+B30", "MW", CALC, "0.00")

section(sz, 33, "STEP 4 — INVERTER CAPABILITY & COMPENSATION", 3)
inv_rows = [
    (34, "Inverter apparent power", 363, "kVA", INPUT, "0"),
    (35, "PF at inverter level  (tip: Goal Seek on B47/B48)", 0.918, "cos φ", INPUT, "0.0000"),
    (36, "Active power / inverter", "=B34*B35", "kW", CALC, "0.0"),
    (37, "Reactive capability — derived √(S²-P²)", "=SQRT(B34^2-B36^2)", "kVAr", CALC, "0.0"),
    (38, "Reactive capability — manual override (optional)", None, "kVAr", INPUT, "0.0"),
    (39, "Reactive capability — from InverterPQ table", "=InverterPQ!B16", "kVAr", CALC, "0.0"),
    (40, "Q per inverter USED", '=IF(InverterPQ!B2="Yes",B39,IF(ISNUMBER(B38),B38,B37))', "kVAr", CALC, "0.0"),
    (41, "Capacitor bank", 0, "MVAr", INPUT, "0.0"),
    (42, "Q for inverters (after cap bank)", "=MAX(0,B22-B41)", "MVAr", CALC, "0.00"),
    (43, "Manual inverter count (blank = auto)", None, "nos.", INPUT, "0"),
    (44, "Required inverters — reactive", "=IF(B40>0,ROUNDUP(B42*1000/B40,0),0)", "nos.", CALC, "0"),
    (45, "Required inverters — active delivery", "=ROUNDUP(B31*1000/B36,0)", "nos.", CALC, "0"),
    (46, "INVERTERS IN DESIGN", "=IF(ISNUMBER(B43),B43,MAX(B44,B45))", "nos.", RESULT, "0"),
    (47, "Power delivered to grid (net of losses)", "=B46*B36/1000-B30", "MW", RESULT, "0.00"),
    (48, "Reactive power compensated", "=B46*B40/1000+B41", "MVAr", RESULT, "0.00"),
    (49, "Delivery check (≥ grid limit)", f'=IF(B47>=B7-{EPS},"PASS","FAIL")', "", RESULT, None),
    (50, "Reactive check (≥ total demand)", f'=IF(B48>=B22-{EPS},"PASS","FAIL")', "", RESULT, None),
]
for r, label, val, unit, fill, num in inv_rows:
    srow(r, label, val, unit, fill, num, bold=(r in (46, 47, 48)))
passfail(sz, "B49:B50")

# =============================== GRIDCODES ===============================
gcs = wb.create_sheet("GridCodes")
gcs.sheet_view.showGridLines = False
widths = {"A": 10, "B": 54, "C": 9, "D": 7, "E": 8, "F": 11, "G": 7, "H": 7,
          "I": 9, "J": 9, "K": 7, "L": 9, "M": 9, "N": 100}
for col, w in widths.items():
    gcs.column_dimensions[col].width = w
title(gcs, "GRID CODE REQUIREMENT DATABASE", 6)
gcs["A3"] = "Selected grid code:"
gcs["A3"].font = BOLD
gcs.merge_cells("B3:F3")
gcs["B3"] = "United Kingdom - ENA G99 / GB Grid Code (Type C/D PPM)"
style(gcs, "B3", INPUT)
headers = ["Key", "Name", "Q/Pmax", "PF", "pMin %", "pFullLag %", "Relief", "Stub",
           "VFullLo", "VFullHi", "Asym", "VZeroLo", "VZeroHi", "Notes"]
table_header(gcs, 5, headers)
codes = [
    ("gb_g99", "United Kingdom - ENA G99 / GB Grid Code (Type C/D PPM)", 0.329, 0.95, 20, 20, 1, 0.05, 0.95, 1.05, 1, 0.90, 1.10,
     "0.95 lead/lag at RC from 20% output (G99 13.5.5, Fig 13.13/13.14). Asymmetric V-Q; lead relief triangle 20-50% at DNO discretion; ±5% band below 20%."),
    ("de_4110", "Germany - VDE-AR-N 4110 (medium voltage)", 0.329, 0.95, 20, 20, 0, 0, 0.95, 1.05, 0, 0.925, 1.075,
     "cos φ 0.95 under/overexcited above 20% of installed capacity; control range 0.925-1.075 pu."),
    ("de_4120", "Germany - VDE-AR-N 4120 (high voltage)", 0.411, 0.925, 20, 20, 0, 0, 0.95, 1.05, 0, 0.90, 1.10,
     "cos φ 0.925 variant (three TSO-selectable variants exist)."),
    ("us_1547b", "USA - IEEE 1547-2018 (Category B)", 0.44, 0.90, 20, 20, 0, 0, 0.95, 1.05, 0, 0.90, 1.10,
     "Inject/absorb 44% of nameplate kVA (~0.90 PF at rated); full capability above 20% of rated power."),
    ("au_ner", "Australia - NER S5.2.5.1 (Automatic access)", 0.395, 0.93, 0, 0, 0, 0, 0.95, 1.05, 0, 0.90, 1.10,
     "±0.395 x rated active power at any output level."),
    ("eu_rfg", "EU - ENTSO-E RfG (typical Type C/D implementation)", 0.329, 0.95, 20, 50, 0, 0, 0.95, 1.05, 0, 0.90, 1.10,
     "Regulation 2016/631 leaves envelopes to each TSO - typical national choice shown."),
    ("pf95", "Generic - fixed PF 0.95 lead/lag", 0.329, 0.95, 20, 20, 0, 0, 0.95, 1.05, 0, 0.90, 1.10, "Fixed PF, full Q above 20% output."),
    ("pf90", "Generic - fixed PF 0.90 lead/lag", 0.484, 0.90, 20, 20, 0, 0, 0.95, 1.05, 0, 0.90, 1.10, "Fixed PF, full Q above 20% output."),
]
for i, row in enumerate(codes):
    r = 6 + i
    for j, v in enumerate(row):
        c = f"{get_column_letter(j+1)}{r}"
        gcs[c] = v
        gcs[c].border = THIN
        if j == 13:
            gcs[c].alignment = Alignment(wrap_text=True, vertical="top")
        elif j >= 2:
            gcs[c].alignment = Alignment(horizontal="center")
gcs.conditional_formatting.add("A6:N13", FormulaRule(formula=["$B6=$B$3"], fill=RESULT))
dv = DataValidation(type="list", formula1="=$B$6:$B$13", allow_blank=False, showDropDown=False)
gcs.add_data_validation(dv)
dv.add(gcs["B3"])
gcs["A16"] = "Active parameters (looked up from the selected code):"
gcs["A16"].font = BOLD
active = ["Q/Pmax", "PF", "pMin %", "pFullLag %", "Relief", "Stub", "VFullLo", "VFullHi", "Asym", "VZeroLo", "VZeroHi"]
for j, name in enumerate(active):
    hc = f"{get_column_letter(j+2)}17"
    vc = f"{get_column_letter(j+2)}18"
    gcs[hc] = name
    gcs[hc].font = BOLD
    gcs[hc].fill = SUBHEAD
    gcs[hc].border = THIN
    src_col = get_column_letter(j + 3)
    gcs[vc] = f"=INDEX(${src_col}$6:${src_col}$13,MATCH($B$3,$B$6:$B$13,0))"
    style(gcs, vc, CALC, center=True)

# =============================== INVERTERPQ ===============================
pq = wb.create_sheet("InverterPQ")
pq.sheet_view.showGridLines = False
pq.column_dimensions["A"].width = 30
for j in range(2, 14):
    pq.column_dimensions[get_column_letter(j)].width = 8
title(pq, "INVERTER P-Q CAPABILITY TABLE", 12)
pq["A2"] = "Use capability table for sizing?"
pq["A2"].font = BOLD
pq["B2"] = "Yes"
style(pq, "B2", INPUT, center=True)
dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False, showDropDown=False)
pq.add_data_validation(dv2)
dv2.add(pq["B2"])
pq["A3"] = "Max reactive power, % of rated kVA. Blank = N/A (operation not permitted). Rows ascending in voltage."
pq["A3"].font = Font(italic=True, size=9)
table_header(pq, 4, ["V (pu) \\ P (pu)"])
for j, p in enumerate((0, 0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0)):
    c = pq.cell(row=4, column=2 + j, value=p)  # numeric — MATCH depends on it
    c.font = BOLD
    c.fill = SUBHEAD
    c.border = THIN
    c.number_format = "0.0"
    c.alignment = Alignment(horizontal="center")
table = [
    (0.85, [60, 60, 60, 60, 60, 60, 60, 48.2, 28.7, None, None]),
    (0.90, [60, 60, 60, 60, 60, 60, 60, 56.6, 41.2, 0, None]),
    (0.95, [60, 60, 60, 60, 60, 60, 60, 60, 51.2, 30.4, None]),
    (1.00, [60, 60, 60, 60, 60, 60, 60, 60, 60, 43.6, 0]),
    (1.05, [60, 60, 60, 60, 60, 60, 60, 60, 60, 43.6, 0]),
    (1.10, [60, 60, 60, 60, 60, 60, 60, 60, 60, 43.6, 0]),
]
for i, (v, vals) in enumerate(table):
    r = 5 + i
    pq[f"A{r}"] = v
    style(pq, f"A{r}", INPUT, bold=True, num="0.00", center=True)
    for j, q in enumerate(vals):
        c = f"{get_column_letter(2+j)}{r}"
        if q is not None:
            pq[c] = q
        style(pq, c, INPUT, center=True)
NOM_ROW = 8  # the 1.00 Vn row


def interp_row(row: int, xcell: str, sheet_prefix: str = "") -> str:
    """Linear interpolation of q% along one capability row at loading X.
    Blank (N/A) cells or loading beyond the last defined column give 0."""
    H = f"{sheet_prefix}$B$4:$L$4"
    R = f"{sheet_prefix}$B${row}:$L${row}"
    j = f"MATCH({xcell},{H},1)"
    y1 = f"INDEX({R},1,{j})"
    y2 = f"INDEX({R},1,{j}+1)"
    x1 = f"INDEX({H},1,{j})"
    x2 = f"INDEX({H},1,{j}+1)"
    return (f'IF({xcell}>=1,IF(ISNUMBER(INDEX({R},1,11)),INDEX({R},1,11),0),'
            f'IF(NOT(ISNUMBER({y1})),0,'
            f'IF({xcell}={x1},{y1},'
            f'IF(NOT(ISNUMBER({y2})),0,'
            f'{y1}+({y2}-{y1})*({xcell}-{x1})/({x2}-{x1})))))')


pq["A13"] = "Sizing lookup (1.00 Vn row at PF loading):"
pq["A13"].font = BOLD
pq["A14"] = "Loading (= inverter PF)"
pq["B14"] = "=Sizing!B35"
style(pq, "B14", CALC, num="0.000")
pq["A15"] = "q% at 1.00 Vn"
pq["B15"] = "=" + interp_row(NOM_ROW, "$B$14")
style(pq, "B15", CALC, num="0.00")
pq["A16"] = "q (kVAr) at 1.00 Vn"
pq["B16"] = "=B15/100*Sizing!B34"
style(pq, "B16", CALC, num="0.0")

# =============================== COMPLIANCE ===============================
cp = wb.create_sheet("Compliance")
cp.sheet_view.showGridLines = False
cw = {"A": 8, "B": 8, "C": 5, "D": 7, "E": 7, "F": 7, "G": 7, "H": 6, "I": 8,
      "J": 10, "K": 10, "L": 10, "M": 10, "N": 10, "O": 9}
for col, w in cw.items():
    cp.column_dimensions[col].width = w
title(cp, "GRID CODE COMPLIANCE", 15)

FL, FH, ZL, ZH, ASYM = "GridCodes!$H$18", "GridCodes!$I$18", "GridCodes!$K$18", "GridCodes!$L$18", "GridCodes!$J$18"
PMIN, PFULL, RELIEF = "GridCodes!$D$18", "GridCodes!$E$18", "GridCodes!$F$18"

section(cp, 3, "OPERATING POINT (Registered Capacity dispatch)", 15)
cp.merge_cells("A4:B4")
cp["A4"] = "Dispatch / inverter"
cp["I4"] = "=Sizing!B31*1000/Sizing!B46"
cp["J4"] = "kW"
style(cp, "I4", CALC, num="0.0")
cp.merge_cells("A5:B5")
cp["A5"] = "Dispatch loading"
cp["I5"] = "=I4/Sizing!B34"
cp["J5"] = "pu"
style(cp, "I5", CALC, num="0.000")
cp.merge_cells("A6:B6")
cp["A6"] = "Qmax required"
cp["I6"] = "=GridCodes!B18*Sizing!B7"
cp["J6"] = "MVAr"
style(cp, "I6", CALC, num="0.00")
for r in (4, 5, 6):
    cp[f"A{r}"].font = BOLD

# ---- capability at table voltages (ascending) ----
section(cp, 8, "CAPABILITY AT TABLE VOLTAGES (per-inverter q% at dispatch loading)", 15)
table_header(cp, 9, ["V (pu)", "q%/inv"])
for i in range(6):
    r = 10 + i
    pr = 5 + i
    cp[f"A{r}"] = f"=InverterPQ!A{pr}"
    style(cp, f"A{r}", CALC, num="0.00", center=True)
    cp[f"B{r}"] = "=" + interp_row(pr, "$I$5", "InverterPQ!")
    style(cp, f"B{r}", CALC, num="0.0", center=True)
ASCV, ASCQ = "$A$10:$A$15", "$B$10:$B$15"

# requirement polygon vertices (chart only)
table_header(cp, 9, ["Req polygon Q", "V"], start_col=13)
poly = [
    ("=0", f"=IF({ASYM}=1,{FL},{ZL})"),
    ("=$I$6", f"={FL}"),
    ("=$I$6", f"=IF({ASYM}=1,1,{FH})"),
    ("=0", f"=IF({ASYM}=1,{FH},{ZH})"),
    ("=-$I$6", f"={FH}"),
    ("=-$I$6", f"=IF({ASYM}=1,1,{FL})"),
    ("=0", f"=IF({ASYM}=1,{FL},{ZL})"),
]
for i, (qf, vf) in enumerate(poly):
    cp[f"M{10+i}"] = qf
    cp[f"N{10+i}"] = vf
    style(cp, f"M{10+i}", CALC, num="0.00")
    style(cp, f"N{10+i}", CALC, num="0.00")


def vq_lag(v):
    sym = (f"IF(OR({v}<={ZL},{v}>={ZH}),0,IF({v}<{FL},({v}-{ZL})/({FL}-{ZL}),"
           f"IF({v}>{FH},({ZH}-{v})/({ZH}-{FH}),1)))")
    asym = f"IF(OR({v}<{FL}-0.0001,{v}>{FH}+0.0001),0,IF({v}<=1,1,({FH}-{v})/({FH}-1)))"
    return f"IF({ASYM}=1,{asym},{sym})"


def vq_lead(v):
    sym = (f"IF(OR({v}<={ZL},{v}>={ZH}),0,IF({v}<{FL},({v}-{ZL})/({FL}-{ZL}),"
           f"IF({v}>{FH},({ZH}-{v})/({ZH}-{FH}),1)))")
    asym = f"IF(OR({v}<{FL}-0.0001,{v}>{FH}+0.0001),0,IF({v}>=1,1,({v}-{FL})/(1-{FL})))"
    return f"IF({ASYM}=1,{asym},{sym})"


# ---- V-Q sweep (interpolated) + corner points ----
section(cp, 17, "V-Q AT REGISTERED CAPACITY — interpolated sweep and A/B/E/F corner points", 15)
vq_headers = ["Point", "V (pu)", "j", "VA", "VB", "qA", "qB", "t", "q%/inv",
              "Q prod (MVAr)", "Q abs (MVAr)", "Req lag (MVAr)", "Req lead (MVAr)", "-Q abs", "Status"]
table_header(cp, 18, vq_headers)
VQ_FIRST = 19
n_sweep = 21
corner_defs = [
    ("A", f"={FH}", "lead"),
    ("B", f"=IF({ASYM}=1,1,{FL})", "lead"),
    ("E", f"={FL}", "lag"),
    ("F", f"=IF({ASYM}=1,1,{FH})", "lag"),
]
for i in range(n_sweep + 4):
    r = VQ_FIRST + i
    if i < n_sweep:
        cp[f"A{r}"] = ""
        cp[f"B{r}"] = round(0.85 + 0.0125 * i, 4)
    else:
        name, vf, dirn = corner_defs[i - n_sweep]
        cp[f"A{r}"] = name
        cp[f"A{r}"].font = BOLD
        cp[f"B{r}"] = vf
    style(cp, f"B{r}", CALC, num="0.000", center=True)
    cp[f"C{r}"] = f"=MAX(1,MIN(5,MATCH(B{r},{ASCV},1)))"
    cp[f"D{r}"] = f"=INDEX({ASCV},C{r})"
    cp[f"E{r}"] = f"=INDEX({ASCV},C{r}+1)"
    cp[f"F{r}"] = f"=INDEX({ASCQ},C{r})"
    cp[f"G{r}"] = f"=INDEX({ASCQ},C{r}+1)"
    cp[f"H{r}"] = f"=MIN(1,MAX(0,(B{r}-D{r})/(E{r}-D{r})))"
    cp[f"I{r}"] = (f"=IF(H{r}<=0.0001,F{r},IF(H{r}>=0.9999,G{r},"
                   f"IF(OR(F{r}<=0,G{r}<=0),0,F{r}+(G{r}-F{r})*H{r})))")
    cp[f"J{r}"] = f"=Sizing!$B$46*I{r}/100*Sizing!$B$34/1000+Sizing!$B$41-Sizing!$B$21"
    cp[f"K{r}"] = f"=Sizing!$B$46*I{r}/100*Sizing!$B$34/1000+Sizing!$B$21"
    cp[f"L{r}"] = f"=$I$6*({vq_lag(f'B{r}')})"
    cp[f"M{r}"] = f"=$I$6*({vq_lead(f'B{r}')})"
    cp[f"N{r}"] = f"=-K{r}"
    if i < n_sweep:
        cp[f"O{r}"] = (f'=IF(AND(MAX(0,J{r})>=L{r}-{EPS},MAX(0,K{r})>=M{r}-{EPS}),"PASS","FAIL")')
    else:
        dirn = corner_defs[i - n_sweep][2]
        cp[f"O{r}"] = (f'=IF(MAX(0,K{r})>=M{r}-{EPS},"PASS","FAIL")' if dirn == "lead"
                       else f'=IF(MAX(0,J{r})>=L{r}-{EPS},"PASS","FAIL")')
    for col, num in (("C", "0"), ("D", "0.00"), ("E", "0.00"), ("F", "0.0"), ("G", "0.0"),
                     ("H", "0.00"), ("I", "0.0"), ("J", "0.00"), ("K", "0.00"),
                     ("L", "0.00"), ("M", "0.00"), ("N", "0.00")):
        style(cp, f"{col}{r}", CALC, num=num, center=True)
    style(cp, f"O{r}", CALC, center=True)
VQ_LAST = VQ_FIRST + n_sweep - 1          # last sweep row
CORNER_LAST = VQ_FIRST + n_sweep + 3
passfail(cp, f"O{VQ_FIRST}:O{CORNER_LAST}")
# hide helper columns
for col in "CDEFGH":
    cp.column_dimensions[col].hidden = True

# ---- P-Q at nominal ----
PQ_HDR = CORNER_LAST + 3
section(cp, PQ_HDR - 1, "P-Q AT NOMINAL VOLTAGE (1.00 pu)", 15)
pq_headers = ["P (pu)", "P (MW)", "", "", "", "", "", "", "q%/inv",
              "Q prod (MVAr)", "Q abs (MVAr)", "Req lag (MVAr)", "Req lead (MVAr)", "-Q abs", "Status"]
table_header(cp, PQ_HDR, pq_headers)
cp[f"C{PQ_HDR}"] = "disp kW"
cp[f"D{PQ_HDR}"] = "loading"
PQ_FIRST = PQ_HDR + 1
n_pq = 41
for i in range(n_pq):
    r = PQ_FIRST + i
    ppu = round(i * 0.025, 3)
    cp[f"A{r}"] = ppu
    style(cp, f"A{r}", CALC, num="0.000", center=True)
    cp[f"B{r}"] = f"=A{r}*Sizing!$B$7"
    cp[f"C{r}"] = f"=A{r}*Sizing!$B$31*1000/Sizing!$B$46"
    cp[f"D{r}"] = f"=C{r}/Sizing!$B$34"
    cp[f"I{r}"] = "=" + interp_row(NOM_ROW, f"$D${r}", "InverterPQ!")
    cp[f"J{r}"] = f"=Sizing!$B$46*I{r}/100*Sizing!$B$34/1000+Sizing!$B$41-Sizing!$B$21"
    cp[f"K{r}"] = f"=Sizing!$B$46*I{r}/100*Sizing!$B$34/1000+Sizing!$B$21"
    ramp = f"IF({PFULL}<={PMIN},1,MIN(1,(A{r}-{PMIN}/100)/({PFULL}/100-{PMIN}/100)))"
    cp[f"L{r}"] = f"=IF(A{r}<{PMIN}/100-0.0001,0,$I$6*{ramp})"
    relief = (f"IF(A{r}>=0.5,$I$6,$I$6*0.12/0.33+($I$6-$I$6*0.12/0.33)*(A{r}-{PMIN}/100)/(0.5-{PMIN}/100))")
    cp[f"M{r}"] = f"=IF(A{r}<{PMIN}/100-0.0001,0,IF({RELIEF}=1,{relief},$I$6*{ramp}))"
    cp[f"N{r}"] = f"=-K{r}"
    cp[f"O{r}"] = f'=IF(AND(MAX(0,J{r})>=L{r}-{EPS},MAX(0,K{r})>=M{r}-{EPS}),"PASS","FAIL")'
    for col, num in (("B", "0.0"), ("C", "0.0"), ("D", "0.000"), ("I", "0.0"),
                     ("J", "0.00"), ("K", "0.00"), ("L", "0.00"), ("M", "0.00"), ("N", "0.00")):
        style(cp, f"{col}{r}", CALC, num=num, center=True)
    style(cp, f"O{r}", CALC, center=True)
PQ_LAST = PQ_FIRST + n_pq - 1
passfail(cp, f"O{PQ_FIRST}:O{PQ_LAST}")
# -Req lead helper for the chart
cp[f"P{PQ_HDR}"] = "-Req lead"
cp[f"P{PQ_HDR}"].font = BOLD
for i in range(n_pq):
    r = PQ_FIRST + i
    cp[f"P{r}"] = f"=-M{r}"
    style(cp, f"P{r}", CALC, num="0.00")
cp.column_dimensions["P"].hidden = True
# MSOL line helper (2 points)
MS = PQ_LAST + 2
cp[f"A{MS}"] = "MSOL"
cp[f"A{MS}"].font = BOLD
cp[f"M{MS}"] = "=-$I$6*1.3"
cp[f"N{MS}"] = f"={PMIN}/100*Sizing!$B$7"
cp[f"M{MS+1}"] = "=$I$6*1.3"
cp[f"N{MS+1}"] = f"={PMIN}/100*Sizing!$B$7"

VERD = MS + 3
cp.merge_cells(f"A{VERD}:B{VERD}")
cp[f"A{VERD}"] = "OVERALL VERDICT"
cp[f"A{VERD}"].font = BOLD
cp.merge_cells(f"I{VERD}:K{VERD}")
cp[f"I{VERD}"] = (f'=IF(COUNTIF(O{VQ_FIRST}:O{CORNER_LAST},"FAIL")'
                  f'+COUNTIF(O{PQ_FIRST}:O{PQ_LAST},"FAIL")=0,"COMPLIANT","NOT COMPLIANT")')
style(cp, f"I{VERD}", RESULT, bold=True, center=True)
cp.conditional_formatting.add(f"I{VERD}", CellIsRule(operator="equal", formula=['"COMPLIANT"'], fill=PASSF))
cp.conditional_formatting.add(f"I{VERD}", CellIsRule(operator="equal", formula=['"NOT COMPLIANT"'], fill=FAILF))

# ---------- charts ----------
INK = "404040"
BLUE = "1C5CAB"
RED = "C0392B"


def scat(xref, yref, title_, color, width=22000, dash=None, smooth=False):
    s = Series(yref, xref, title=title_)
    s.marker = Marker(symbol="none")
    s.smooth = smooth
    lp = LineProperties(solidFill=color, w=width)
    if dash:
        lp.prstDash = dash
    s.graphicalProperties.line = lp
    return s


vq_chart = ScatterChart()
vq_chart.title = "V-Q Diagram at CP — Registered Capacity"
vq_chart.style = 13
vq_chart.x_axis.title = "Q at CP (MVAr)    [lead < 0 < lag]"
vq_chart.y_axis.title = "V (pu)"
vq_chart.y_axis.scaling.min = 0.84
vq_chart.y_axis.scaling.max = 1.12
vq_chart.x_axis.majorGridlines = None
vq_chart.height, vq_chart.width = 11, 18
vq_chart.series.append(scat(Reference(cp, min_col=13, min_row=10, max_row=16),
                            Reference(cp, min_col=14, min_row=10, max_row=16),
                            "Requirement", INK, dash="dash"))
vq_chart.series.append(scat(Reference(cp, min_col=10, min_row=VQ_FIRST, max_row=VQ_LAST),
                            Reference(cp, min_col=2, min_row=VQ_FIRST, max_row=VQ_LAST),
                            "Capability — production (lag)", BLUE))
vq_chart.series.append(scat(Reference(cp, min_col=14, min_row=VQ_FIRST, max_row=VQ_LAST),
                            Reference(cp, min_col=2, min_row=VQ_FIRST, max_row=VQ_LAST),
                            "Capability — absorption (lead)", RED))
vq_chart.legend.position = "b"
cp.add_chart(vq_chart, "Q3")

pq_chart = ScatterChart()
pq_chart.title = "P-Q Diagram at CP — 1.00 pu"
pq_chart.style = 13
pq_chart.x_axis.title = "Q at CP (MVAr)    [lead < 0 < lag]"
pq_chart.y_axis.title = "P (MW)"
pq_chart.y_axis.scaling.min = 0
pq_chart.x_axis.majorGridlines = None
pq_chart.height, pq_chart.width = 11, 18
pq_chart.series.append(scat(Reference(cp, min_col=12, min_row=PQ_FIRST, max_row=PQ_LAST),
                            Reference(cp, min_col=2, min_row=PQ_FIRST, max_row=PQ_LAST),
                            "Requirement (lag)", INK, dash="dash"))
pq_chart.series.append(scat(Reference(cp, min_col=16, min_row=PQ_FIRST, max_row=PQ_LAST),
                            Reference(cp, min_col=2, min_row=PQ_FIRST, max_row=PQ_LAST),
                            "Requirement (lead)", INK, dash="dash"))
pq_chart.series.append(scat(Reference(cp, min_col=10, min_row=PQ_FIRST, max_row=PQ_LAST),
                            Reference(cp, min_col=2, min_row=PQ_FIRST, max_row=PQ_LAST),
                            "Capability — production (lag)", BLUE))
pq_chart.series.append(scat(Reference(cp, min_col=14, min_row=PQ_FIRST, max_row=PQ_LAST),
                            Reference(cp, min_col=2, min_row=PQ_FIRST, max_row=PQ_LAST),
                            "Capability — absorption (lead)", RED))
pq_chart.series.append(scat(Reference(cp, min_col=13, min_row=MS, max_row=MS + 1),
                            Reference(cp, min_col=14, min_row=MS, max_row=MS + 1),
                            "MSOL", "888888", width=12000, dash="sysDot"))
pq_chart.legend.position = "b"
cp.add_chart(pq_chart, "Q26")

out = "Reactive_Power_Compensation_Tool.xlsx"
wb.save(out)
print("saved", out)
